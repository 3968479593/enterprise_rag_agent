"""文档处理器：把上传的原始字节解析为纯文本。

支持 txt / md / pdf / docx / csv / json，统一返回 (text, source_type)。
解析失败抛出 ValueError，由上层转为 400/500 响应。
"""
import csv
import io
import json
from pathlib import Path

from app.utils.logger import logger


def decode_text(raw: bytes) -> str:
    """优先 UTF-8，兼容 GBK 等中文编码。"""
    for encoding in ("utf-8", "gb18030", "utf-16"):
        try:
            return raw.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return raw.decode("utf-8", errors="ignore")


def parse_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            pages.append(text.strip())
    return "\n\n".join(pages)


def parse_docx(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_csv(raw: bytes) -> str:
    text = decode_text(raw)
    reader = csv.reader(io.StringIO(text))
    lines = []
    for row in reader:
        if any(cell.strip() for cell in row):
            lines.append(", ".join(cell.strip() for cell in row))
    return "\n".join(lines)


def parse_json(raw: bytes) -> str:
    text = decode_text(raw)
    data = json.loads(text)  # 非法 JSON 会抛异常，由上层处理
    return json.dumps(data, ensure_ascii=False, indent=2)


def parse_file(filename: str, raw: bytes) -> tuple[str, str]:
    """返回 (文本内容, 来源类型)。"""
    ext = Path(filename).suffix.lower()
    if ext in (".txt", ".md"):
        return decode_text(raw), ("md" if ext == ".md" else "txt")
    if ext == ".pdf":
        return parse_pdf(raw), "pdf"
    if ext == ".docx":
        return parse_docx(raw), "docx"
    if ext == ".csv":
        return parse_csv(raw), "csv"
    if ext == ".json":
        return parse_json(raw), "json"
    raise ValueError(f"不支持的文件类型: {ext}")


def safe_filename(filename: str) -> str:
    """清洗上传文件名，防止路径穿越与特殊字符。"""
    name = Path(filename or "unnamed.txt").name
    name = "".join(ch for ch in name if ch not in '<>:"/\\|?*')
    return name.strip() or "unnamed.txt"
